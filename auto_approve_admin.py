"""
FASIH Survey Automation - Approve Admin
========================================
Script otomasi Playwright untuk melakukan Approve Admin di FASIH Web (BPS).

Alur Kerja:
  1. Browser Chrome terbuka otomatis dengan profil tersimpan (headless=False).
  2. User login manual (jika belum login), lalu tekan ENTER di terminal.
  3. Script membaca data dari file Excel (data_edited.xlsx atau input_fasih_ntb/data_edited.xlsx):
     - Membuka link assignment (01_bukalink.html)
     - Menunggu proses loading selesai (FormGear & FASIH overlay)
     - Mencari dan mengklik tombol Approve (02_btnapprove.html - icon centang hijau)
     - Menunggu modal konfirmasi muncul (03_modalkonfirmasi.html - 'Approve Assignment?')
     - Mengklik tombol 'Konfirmasi' di dalam modal
     - Menunggu proses penyimpanan selesai dan verifikasi status
     - Mengupdate status di file Excel menjadi 'BERHASIL' dan mencatat timestamp
     - Menulis ke log_simple_approve.txt dan log_terakhir_approve.txt
  4. Lanjut ke assignment berikutnya hingga selesai.

Cara Pakai:
  python auto_approve_admin.py
  python auto_approve_admin.py --dry-run
  python auto_approve_admin.py --excel input_fasih_ntb/data_edited.xlsx
  python auto_approve_admin.py --limit 10
  python auto_approve_admin.py --start-row 5
"""

import asyncio
import logging
import os
import sys
import time
import shutil
import subprocess
import argparse
from datetime import datetime
from pathlib import Path
from playwright.async_api import (
    async_playwright,
    Page,
    BrowserContext,
    TimeoutError as PlaywrightTimeout,
)
import openpyxl
from playwright_stealth.stealth import Stealth

# Pastikan console Windows / macOS mendukung output UTF-8 (emoticon / simbol)
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# ─────────────────────────── KONFIGURASI DEFAULT ───────────────────────────

BASE_URL       = "https://fasih-sm.bps.go.id"
DEFAULT_EXCEL  = "data_edited.xlsx"
FALLBACK_EXCEL = "input_fasih_ntb/data_edited.xlsx"
AUTH_FILE      = "auth.json"
SLOW_MO        = 250        # ms jeda antar aksi
TIMEOUT        = 35_000     # ms timeout default
ASSIGNMENT_ROW_TIMEOUT = 90 # Batas waktu maksimal per assignment (detik)

LOG_FILE       = "log_approve.txt"
LAST_ROW_TXT   = "log_terakhir_approve.txt"
LOG_SIMPLE_TXT = "log_simple_approve.txt"

# Konfigurasi logging: Detail lengkap ke file, console tetap bersih & simpel
log = logging.getLogger("fasih_approve")
log.setLevel(logging.INFO)
log.handlers.clear()

file_handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
file_handler.setLevel(logging.INFO)
file_handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S"))
log.addHandler(file_handler)

# Console handler hanya menampilkan Warning dan Error agar tidak spam di terminal
console_handler = logging.StreamHandler()
console_handler.setLevel(logging.WARNING)
console_handler.setFormatter(logging.Formatter("\033[93m[%(levelname)s]\033[0m %(message)s"))
log.addHandler(console_handler)


# ─────────────────────────── ANIMASI LOADING TERMINAL (####) ──────────

class HashLoader:
    """
    Animasi loading teks '####' dinamis di terminal untuk memberikan visual progres yang menarik.
    """
    FRAMES = [
        "[#         ]",
        "[##        ]",
        "[###       ]",
        "[####      ]",
        "[#####     ]",
        "[######    ]",
        "[#######   ]",
        "[########  ]",
        "[######### ]",
        "[##########]",
        "[ #########]",
        "[  ########]",
        "[   #######]",
        "[    ######]",
        "[     #####]",
        "[      ####]",
        "[       ###]",
        "[        ##]",
        "[         #]",
        "[          ]",
    ]

    def __init__(self, message: str = "Memproses...", prefix: str = ""):
        self.message = message
        self.prefix = prefix
        self._running = False
        self._task = None

    async def _animate(self):
        idx = 0
        while self._running:
            f = self.FRAMES[idx % len(self.FRAMES)]
            prefix_str = f"{self.prefix} " if self.prefix else ""
            sys.stdout.write(f"\r{prefix_str}\033[96m{f}\033[0m {self.message}\033[K")
            sys.stdout.flush()
            idx += 1
            await asyncio.sleep(0.08)

    def start(self, message: str = None, prefix: str = ""):
        if message:
            self.message = message
        if prefix:
            self.prefix = prefix
        self._running = True
        self._task = asyncio.create_task(self._animate())

    def update(self, new_msg: str):
        self.message = new_msg
        log.info(new_msg)

    async def stop(self, clear: bool = True):
        self._running = False
        if self._task:
            self._task.cancel()
            try:
                await self._task
            except asyncio.CancelledError:
                pass
            self._task = None
        if clear:
            sys.stdout.write("\r\033[K")
            sys.stdout.flush()


def format_progress(current: int, total: int, width: int = 10) -> str:
    """
    Format progress bar dengan format persentase dan baris:
    Contoh: [#######---] 70% (7/10 baris)
    """
    if total <= 0:
        pct = 0.0
        filled = 0
    else:
        pct = (current / total) * 100
        filled = int(round(width * current / total))
    filled = max(0, min(width, filled))
    bar = "#" * filled + "-" * (width - filled)
    return f"[{bar}] {pct:.0f}% ({current}/{total} baris)"


# ─────────────────────────── LOGGING HELPERS ───────────────────────────

def append_log_simple(item: dict, status: str, timestamp: str = None):
    """
    Menulis / menambahkan satu baris log ringkas ke file log_simple_approve.txt:
    id | link_fasih | nama usaha | status | timestamp
    """
    try:
        ts = timestamp or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        st_clean = "sukses" if "BERHASIL" in status.upper() or "SUKSES" in status.upper() or "APPROVED" in status.upper() else "tidak"

        id_val = str(item.get("id_assignment") or item.get("row_idx") or "").strip()
        link_val = str(item.get("link_fasih") or "").strip()
        nama_usaha_val = str(item.get("nama_usaha") or item.get("nama_assignment") or "").strip()

        clean = lambda s: str(s).replace("\n", " ").replace("\r", "").replace("|", "/")

        line = (
            f"{clean(id_val)} | "
            f"{clean(link_val)} | "
            f"{clean(nama_usaha_val)} | "
            f"{st_clean} | "
            f"{ts}\n"
        )

        header = "id | link_fasih | nama usaha / assignment | status | timestamp\n"

        file_exists = os.path.exists(LOG_SIMPLE_TXT) and os.path.getsize(LOG_SIMPLE_TXT) > 0
        with open(LOG_SIMPLE_TXT, "a", encoding="utf-8") as f:
            if not file_exists:
                f.write(header)
            f.write(line)
    except Exception as e:
        log.warning(f"Gagal menulis ke {LOG_SIMPLE_TXT}: {e}")


def _normalize_key(text: str) -> str:
    """Normalisasi string key ID untuk lookup status di log."""
    if not text:
        return ""
    return str(text).strip().lower()


def load_log_simple_status(filepath: str = LOG_SIMPLE_TXT) -> dict[str, str]:
    """
    Membaca log_simple_approve.txt untuk melanjutkan proses (resume).
    Return dict dengan key (id_assignment) -> status ('sukses' atau 'tidak').
    """
    status_map = {}
    if not os.path.exists(filepath):
        return status_map

    try:
        with open(filepath, "r", encoding="utf-8") as f:
            lines = f.readlines()

        for line in lines[1:]:  # Lewati baris header
            line = line.strip()
            if not line or "|" not in line:
                continue
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 4:
                id_val = _normalize_key(parts[0])
                status = parts[3].strip().lower()
                if id_val:
                    status_map[id_val] = status
    except Exception as e:
        log.warning(f"Gagal membaca riwayat {filepath}: {e}")

    return status_map


def update_last_row_txt(item: dict, status: str, note: str = "", current_idx: int = 0, total_items: int = 0, sukses: int = 0, gagal: int = 0):
    """
    Menyimpan snapshot proses baris terakhir ke log_terakhir_approve.txt.
    """
    try:
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        pct = (current_idx / total_items * 100) if total_items > 0 else 0

        content = (
            f"╔══════════════════════════════════════════════════════════════════════╗\n"
            f"║          PROGRESS STATUS APPROVE ADMIN FASIH (REAL-TIME)             ║\n"
            f"╠══════════════════════════════════════════════════════════════════════╣\n"
            f"  Waktu Terakhir    : {now_str}\n"
            f"  Progress Total    : {current_idx}/{total_items} data ({pct:.1f}%)\n"
            f"  Statistik         : {sukses} Berhasil | {gagal} Gagal\n"
            f"  Baris Excel       : Baris {item.get('row_idx')}\n"
            f"  ID Assignment     : {item.get('id_assignment', '-')}\n"
            f"  Nama Usaha        : {item.get('nama_usaha') or item.get('nama_assignment', '-')}\n"
            f"  URL               : {item.get('link_fasih', '-')}\n"
            f"  Status Terakhir   : {status}\n"
            f"  Keterangan        : {note if note else '-'}\n"
            f"╚══════════════════════════════════════════════════════════════════════╝\n"
        )

        with open(LAST_ROW_TXT, "w", encoding="utf-8") as f:
            f.write(content)
    except Exception as e:
        log.warning(f"Gagal menulis {LAST_ROW_TXT}: {e}")


# ─────────────────────────── SYSTEM & BROWSER HELPERS ───────────────────────────

def get_chrome_executable() -> str:
    """Mencari path executable Google Chrome secara otomatis (macOS, Windows, Linux)."""
    if sys.platform == "darwin":
        candidates = [
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            os.path.expanduser("~/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        ]
    elif sys.platform == "win32":
        candidates = [
            os.path.expandvars(r"%ProgramFiles%\Google\Chrome\Application\chrome.exe"),
            os.path.expandvars(r"%ProgramFiles(x86)%\Google\Chrome\Application\chrome.exe"),
            os.path.expandvars(r"%LocalAppData%\Google\Chrome\Application\chrome.exe"),
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        ]
    else:  # Linux / Unix
        candidates = [
            "/usr/bin/google-chrome",
            "/usr/bin/google-chrome-stable",
            "/usr/bin/chromium",
            "/usr/bin/chromium-browser",
            "/snap/bin/chromium",
        ]

    for p in candidates:
        if p and os.path.isfile(p):
            return p

    for name in ["google-chrome", "chrome", "google-chrome-stable", "chromium"]:
        found = shutil.which(name)
        if found:
            return found

    if sys.platform == "win32":
        return r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    elif sys.platform == "darwin":
        return "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    return "google-chrome"


def get_screen_bounds() -> tuple[int, int]:
    """Mendapatkan resolusi layar aktif lintas platform."""
    if sys.platform == "win32":
        try:
            import ctypes
            user32 = ctypes.windll.user32
            try:
                user32.SetProcessDPIAware()
            except Exception:
                pass
            w = user32.GetSystemMetrics(0)
            h = user32.GetSystemMetrics(1)
            if w > 800 and h > 600:
                return w, h
        except Exception:
            pass

    if sys.platform == "darwin":
        try:
            out = subprocess.check_output(
                ["osascript", "-e", 'tell application "Finder" to get bounds of window of desktop'],
                timeout=3
            ).decode().strip()
            parts = [int(p.strip()) for p in out.split(",")]
            w = parts[2] - parts[0]
            h = parts[3] - parts[1]
            if w > 800 and h > 600:
                return w, h
        except Exception:
            pass

    try:
        import tkinter
        root = tkinter.Tk()
        root.withdraw()
        w = root.winfo_screenwidth()
        h = root.winfo_screenheight()
        root.destroy()
        if w > 800 and h > 600:
            return w, h
    except Exception:
        pass

    return (1920, 1080) if sys.platform == "win32" else (1680, 1050)


async def maximize_chrome_window(page: Page = None, width: int = 1920, height: int = 1080):
    """Memaksimalkan jendela Google Chrome via CDP."""
    if page:
        try:
            session = await page.context.new_cdp_session(page)
            window_info = await session.send("Browser.getWindowForTarget")
            window_id = window_info.get("windowId")
            if window_id:
                await session.send("Browser.setWindowBounds", {
                    "windowId": window_id,
                    "bounds": {"windowState": "maximized"}
                })
                return
        except Exception:
            pass

    if sys.platform == "darwin":
        try:
            script = f'tell application "Google Chrome" to set bounds of front window to {{0, 0, {width}, {height}}}'
            subprocess.run(["osascript", "-e", script], timeout=3, check=False)
        except Exception:
            pass


def start_anti_sleep():
    """Mencegah PC masuk mode tidur selama otomasi berlangsung."""
    if sys.platform == "darwin":
        try:
            proc = subprocess.Popen(["caffeinate", "-dims"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            return proc
        except Exception:
            return None
    elif sys.platform == "win32":
        try:
            import ctypes
            ES_CONTINUOUS       = 0x80000000
            ES_SYSTEM_REQUIRED  = 0x00000001
            ES_DISPLAY_REQUIRED = 0x00000002
            ctypes.windll.kernel32.SetThreadExecutionState(
                ES_CONTINUOUS | ES_SYSTEM_REQUIRED | ES_DISPLAY_REQUIRED
            )
            return True
        except Exception:
            return None
    return None


def stop_anti_sleep(anti_sleep_handle):
    """Mengembalikan setting sleep ke normal."""
    if sys.platform == "darwin" and anti_sleep_handle:
        try:
            anti_sleep_handle.terminate()
        except Exception:
            pass
    elif sys.platform == "win32":
        try:
            import ctypes
            ES_CONTINUOUS = 0x80000000
            ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS)
        except Exception:
            pass


# ─────────────────────────── EXCEL HELPERS ───────────────────────────

def format_file_size(num_bytes: int) -> str:
    """Format ukuran file menjadi string yang mudah dibaca (KB / MB)."""
    for unit in ['B', 'KB', 'MB', 'GB']:
        if num_bytes < 1024.0:
            return f"{num_bytes:.1f} {unit}"
        num_bytes /= 1024.0
    return f"{num_bytes:.1f} TB"


def scan_excel_files(base_dir: Path = Path(".")) -> list[Path]:
    """Mencari semua file Excel (.xlsx / .xls) di folder kerja dan subfolder tingkat 1."""
    files = []
    # 1. Cari di root folder
    for ext in ("*.xlsx", "*.xls"):
        for p in base_dir.glob(ext):
            if not p.name.startswith("~$") and not p.name.startswith("."):
                files.append(p)
    # 2. Cari di subfolder tingkat 1 (misal input_fasih_ntb/)
    for sub in base_dir.iterdir():
        if sub.is_dir() and not sub.name.startswith(".") and sub.name != "__pycache__":
            for ext in ("*.xlsx", "*.xls"):
                for p in sub.glob(ext):
                    if not p.name.startswith("~$") and not p.name.startswith("."):
                        files.append(p)

    # Urutkan: file di root terlebih dahulu, lalu alfabet
    return sorted(list(set(files)), key=lambda x: (len(x.parts), str(x).lower()))


def select_excel_interactive(cli_excel: str = None) -> Path:
    """
    Menampilkan menu pilihan file Excel interaktif di awal terminal jika ada banyak file Excel
    atau jika user ingin memilih file Excel tertentu secara fleksibel dengan nama yang berbeda-beda.
    """
    if cli_excel:
        p = Path(cli_excel)
        if p.exists():
            return p
        print(f"\033[93m⚠️ File Excel '{cli_excel}' yang ditentukan via argumen tidak ditemukan.\033[0m")
        print("Silakan pilih file dari daftar di bawah ini:\n")

    files = scan_excel_files()

    print("\n" + "═" * 65)
    print("📂 PILIH FILE EXCEL DATA ASSIGNMENT:")
    print("═" * 65)

    if not files:
        print("\033[93m⚠️ Tidak ditemukan file Excel (.xlsx / .xls) otomatis di folder ini.\033[0m")
        while True:
            try:
                custom = input("Ketik nama / path file Excel yang ingin dijalankan: ").strip().strip("'\"")
            except (KeyboardInterrupt, EOFError):
                print("\nDibatalkan.")
                sys.exit(0)
            if custom and Path(custom).exists():
                return Path(custom)
            print(f"\033[91m❌ File '{custom}' tidak ditemukan! Silakan periksa kembali.\033[0m")

    for idx, f in enumerate(files, 1):
        try:
            sz_str = format_file_size(f.stat().st_size)
        except Exception:
            sz_str = ""
        sz_display = f"({sz_str})" if sz_str else ""
        print(f"  [{idx}] \033[96m{f}\033[0m {sz_display}")

    custom_opt = len(files) + 1
    print(f"  [{custom_opt}] ✍️  Ketik nama / path file lain secara manual...")
    print("─" * 65)

    while True:
        prompt_txt = f"Pilih nomor file Excel [1-{custom_opt}] (Tekan ENTER untuk [1]): "
        try:
            choice = input(prompt_txt).strip()
        except (KeyboardInterrupt, EOFError):
            print("\nDibatalkan.")
            sys.exit(0)

        if choice == "":
            selected = files[0]
            print(f"👉 Menggunakan file: \033[92m{selected}\033[0m\n")
            return selected

        if choice.isdigit():
            c_int = int(choice)
            if 1 <= c_int <= len(files):
                selected = files[c_int - 1]
                print(f"👉 File terpilih: \033[92m{selected}\033[0m\n")
                return selected
            elif c_int == custom_opt:
                while True:
                    try:
                        manual = input("Ketik nama / path file Excel: ").strip().strip("'\"")
                    except (KeyboardInterrupt, EOFError):
                        print("\nDibatalkan.")
                        sys.exit(0)
                    if manual and Path(manual).exists():
                        print(f"👉 File terpilih: \033[92m{manual}\033[0m\n")
                        return Path(manual)
                    print(f"\033[91m❌ File '{manual}' tidak ditemukan. Silakan periksa kembali.\033[0m")

        # Jika user langsung mengetik nama file (misal data_september.xlsx)
        direct_path = Path(choice.strip("'\""))
        if direct_path.exists() and (choice.endswith(".xlsx") or choice.endswith(".xls")):
            print(f"👉 File terpilih: \033[92m{direct_path}\033[0m\n")
            return direct_path

        print(f"\033[91m❌ Pilihan '{choice}' tidak valid. Masukkan angka antara 1 sampai {custom_opt}.\033[0m")


def find_default_excel() -> Path:
    """Mendeteksi file Excel default yang ada di workspace."""
    if Path(DEFAULT_EXCEL).exists():
        return Path(DEFAULT_EXCEL)
    if Path(FALLBACK_EXCEL).exists():
        return Path(FALLBACK_EXCEL)
    return Path(DEFAULT_EXCEL)


def load_assignments(filepath: Path) -> tuple[openpyxl.Workbook, list[dict]]:
    """
    Membaca data assignment dari Excel dan memetakan kolom secara dinamis.
    Mendukung skema root data_edited.xlsx maupun input_fasih_ntb/data_edited.xlsx.
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    log.info(f"Membaca file Excel: '{filepath}' ({ws.max_row - 1} baris data)")
    log.info(f"Header kolom terdeteksi: {[h for h in headers if h]}")

    col_map = {}
    for idx, h in enumerate(headers, 1):
        if not h:
            continue
        h_str = str(h).strip().lower()
        if "id_assignment" in h_str or h_str == "id":
            col_map["id_assignment"] = idx
        elif "link" in h_str or "url" in h_str:
            col_map["link_fasih"] = idx
        elif "nama usaha" in h_str:
            col_map["nama_usaha"] = idx
        elif "nama assignment" in h_str:
            col_map["nama_assignment"] = idx
        elif "bot_runner" in h_str or "bot" in h_str:
            col_map["bot_runner"] = idx
        elif "status_approve" in h_str or "status approve" in h_str:
            col_map["status_approve"] = idx
        elif "status_sukses" in h_str or "status sukses" in h_str:
            col_map["status_sukses"] = idx
        elif "status otomasi" in h_str:
            col_map["status_otomasi"] = idx
        elif "status" in h_str:
            col_map["status_general"] = idx
        elif "waktu_approve" in h_str or "waktu approve" in h_str:
            col_map["waktu_approve"] = idx
        elif "timestamp" in h_str:
            col_map["timestamp"] = idx
        elif "waktu eksekusi" in h_str:
            col_map["waktu_eksekusi"] = idx
        elif "waktu" in h_str:
            col_map["waktu_general"] = idx

    # Tentukan kolom Status untuk Approval
    if "status_approve" in col_map:
        status_col = col_map["status_approve"]
    elif "status_sukses" in col_map:
        status_col = col_map["status_sukses"]
    elif "status_general" in col_map:
        status_col = col_map["status_general"]
    else:
        # Jika belum ada kolom khusus status approve, tambahkan kolom baru
        status_col = ws.max_column + 1
        ws.cell(row=1, column=status_col, value="Status Approve")
        col_map["status_approve"] = status_col

    # Tentukan kolom Timestamp untuk Approval
    if "waktu_approve" in col_map:
        time_col = col_map["waktu_approve"]
    elif "timestamp" in col_map:
        time_col = col_map["timestamp"]
    elif "waktu_general" in col_map:
        time_col = col_map["waktu_general"]
    else:
        time_col = status_col + 1
        ws.cell(row=1, column=time_col, value="Waktu Approve")
        col_map["waktu_approve"] = time_col

    bot_runner_col = col_map.get("bot_runner")

    def get_val(r: int, key: str, default=""):
        c = col_map.get(key)
        if c is not None:
            v = ws.cell(r, c).value
            return "" if v is None else v
        return default

    items = []
    for r in range(2, ws.max_row + 1):
        link = get_val(r, "link_fasih")
        if not link or not str(link).strip().startswith("http"):
            continue

        st_val = ws.cell(row=r, column=status_col).value or ""
        st_otomasi = get_val(r, "status_otomasi")

        item = {
            "row_idx": r,
            "id_assignment": str(get_val(r, "id_assignment")).strip(),
            "link_fasih": str(link).strip(),
            "nama_usaha": str(get_val(r, "nama_usaha")).strip(),
            "nama_assignment": str(get_val(r, "nama_assignment")).strip(),
            "status_otomasi": str(st_otomasi).strip(),
            "status": str(st_val).strip(),
            "status_col": status_col,
            "time_col": time_col,
            "bot_runner_col": bot_runner_col,
        }
        items.append(item)

    return wb, items


def save_assignment_status(wb: openpyxl.Workbook, filepath: Path, item: dict, status: str, note: str = ""):
    """Update status approval per baris dan simpan ke file Excel."""
    ws = wb.active
    r = item["row_idx"]
    val = status if not note else f"{status}: {note}"
    ws.cell(row=r, column=item["status_col"], value=val)
    ws.cell(row=r, column=item["time_col"], value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    if item.get("bot_runner_col"):
        ws.cell(row=r, column=item["bot_runner_col"], value="APPROVE_ADMIN")

    try:
        wb.save(filepath)
    except Exception as e:
        log.warning(f"Gagal menyimpan ke file Excel '{filepath}': {e}")


# ─────────────────────────── FORM INTERACTION & LOADING ───────────────────────────

LOADING_SELECTORS = [
    "[class*='f:z-100']",
    "[class*='f:bg-black/20']",
    "[class*='bg-black/20'][class*='items-center']",
    "#FormGear-loader:not(.hidden)",
    "ngx-spinner:not([style*='display: none'])",
    "div.ngx-spinner-overlay",
    "div.swal2-loader",
    "div.swal2-loading",
    "[aria-busy='true']",
    ".spinner-border:visible",
]

APPROVE_BUTTON_SELECTORS = [
    # Selector tombol Approve (02_btnapprove.html - hijau dengan icon check)
    "button.f\\:bg-success:has(svg.tabler-icon-check)",
    "button:has(svg.tabler-icon-check)",
    "button.f\\:bg-success",
    "div.f\\:fixed button:has(svg.tabler-icon-check)",
    "button[data-state]:has(svg.tabler-icon-check)",
    "button:has(path[d*='M5 12l5 5l10 -10'])",
    "xpath=//button[contains(@class, 'f:bg-success') and .//*[name()='svg' and contains(@class, 'tabler-icon-check')]]",
]

REVOKE_BUTTON_SELECTORS = [
    # Selector tombol Revoke (04_revoke.html - icon putar searah jarum jam / rotate-clockwise)
    # Jika tombol ini muncul dan tombol Approve tidak ada, berarti dokumen SUDAH di-approve admin.
    "button:has(svg.tabler-icon-rotate-clockwise)",
    "svg.tabler-icon-rotate-clockwise",
    "button:has(path[d*='M4.05 11'])",
    "div.f\\:fixed button:has(svg.tabler-icon-rotate-clockwise)",
    "div.f\\:flex-col button:has(svg.tabler-icon-rotate-clockwise)",
    "xpath=//button[.//*[name()='svg' and contains(@class, 'tabler-icon-rotate-clockwise')]]",
]

MODAL_CONFIRM_SELECTORS = [
    # Selector tombol Konfirmasi pada Modal (03_modalkonfirmasi.html)
    "div[role='alertdialog']:has-text('Approve') button:has-text('Konfirmasi')",
    "div[role='alertdialog'] button.f\\:bg-success:has-text('Konfirmasi')",
    "xpath=//div[@role='alertdialog']//button[normalize-space()='Konfirmasi']",
    "xpath=//div[@role='alertdialog']//button[contains(@class, 'f:bg-success')]",
    "div[role='alertdialog'] button:has-text('Konfirmasi')",
]


async def wait_for_loading(page: Page, timeout_ms: int = 45_000, initial_delay_ms: int = 300, loader: HashLoader = None) -> None:
    """Menunggu hingga seluruh indikator loading / spinner / overlay di FASIH selesai."""
    if initial_delay_ms > 0:
        await page.wait_for_timeout(initial_delay_ms)

    start_time = time.time()
    max_wait = timeout_ms / 1000.0

    while (time.time() - start_time) < max_wait:
        is_loading = False
        detected_sel = None

        for sel in LOADING_SELECTORS:
            try:
                loc = page.locator(sel).first
                if await loc.is_visible():
                    is_loading = True
                    detected_sel = sel
                    break
            except Exception:
                continue

        if not is_loading:
            return

        log.info(f"Masih loading ({detected_sel}), menunggu hingga selesai...")
        if loader:
            loader.update("Menunggu loading form FASIH selesai...")

        try:
            await page.locator(detected_sel).first.wait_for(state="hidden", timeout=8_000)
        except Exception:
            pass

        await page.wait_for_timeout(400)

    log.warning("Waktu tunggu loading melebihi batas, mencoba melanjutkan...")


async def click_approve_button(page: Page, loader: HashLoader = None) -> tuple[bool, str]:
    """
    Mencari dan mengklik tombol Approve (02_btnapprove.html - hijau dengan icon check).
    Jika tombol Approve TIDAK ditemukan tetapi ada tombol Revoke (04_revoke.html - icon rotate-clockwise),
    maka assignment tersebut SUDAH di-approve admin sebelumnya -> langsung di-skip dan ditandai SUDAH_APPROVED.
    """
    if loader:
        loader.update("Memeriksa status halaman & loading...")
    await wait_for_loading(page, timeout_ms=30_000, initial_delay_ms=400, loader=loader)

    if loader:
        loader.update("Mencari tombol Approve / Revoke...")
    log.info("Mencari tombol Approve di floating action bar...")

    # Cek apakah floating action bar dalam mode collapsed (mobile toggle icon menu-2)
    try:
        mobile_menu = page.locator("button:has(svg.tabler-icon-menu-2)").first
        if await mobile_menu.is_visible(timeout=1_500):
            # Jika desktop buttons tersembunyi, klik menu toggle
            desktop_bar = page.locator("div.f\\:flex-col:has(svg.tabler-icon-check), div.f\\:flex-col:has(svg.tabler-icon-rotate-clockwise)").first
            if not await desktop_bar.is_visible(timeout=500):
                await mobile_menu.click()
                await page.wait_for_timeout(500)
    except Exception:
        pass

    # Tunggu selector tombol aksi (Approve, Edit, atau Revoke) muncul
    try:
        await page.wait_for_selector(
            "button:has(svg.tabler-icon-check), button:has(svg.tabler-icon-rotate-clockwise), button:has(svg.tabler-icon-edit)",
            state="visible",
            timeout=8_000,
        )
    except Exception:
        pass

    # 1. Cari tombol Approve menggunakan daftar selector
    for sel in APPROVE_BUTTON_SELECTORS:
        try:
            btn = page.locator(sel).first
            if await btn.count() > 0 and await btn.is_visible():
                # Cek jika tombol disabled
                if await btn.is_disabled():
                    log.warning("Tombol Approve berstatus DISABLED.")
                    return False, "Tombol Approve disabled"

                if loader:
                    loader.update("Mengklik tombol Approve...")
                await btn.scroll_into_view_if_needed()
                try:
                    await btn.click(timeout=4_000)
                except Exception:
                    await btn.click(force=True)

                log.info("Tombol Approve berhasil diklik.")
                return True, "OK"
        except Exception:
            continue

    # 2. JIKA TOMBOL APPROVE TIDAK KETEMU: Cek apakah ada tombol Revoke (04_revoke.html)
    # Keberadaan tombol Revoke menandakan assignment ini SUDAH di-approve admin sebelumnya.
    for r_sel in REVOKE_BUTTON_SELECTORS:
        try:
            revoke_btn = page.locator(r_sel).first
            if await revoke_btn.count() > 0 and await revoke_btn.is_visible(timeout=1_000):
                log.info("Tombol Revoke (04_revoke.html) terdeteksi. Assignment sudah di-approve admin.")
                return True, "SUDAH_APPROVED"
        except Exception:
            continue

    # 3. Cek teks status di halaman (fallback)
    try:
        body_text = await page.inner_text("body")
        if "approved" in body_text.lower() or "disetujui" in body_text.lower():
            log.info("Status 'Approved' / 'Disetujui' terdeteksi di halaman.")
            return True, "SUDAH_APPROVED"
    except Exception:
        pass

    log.error("Tombol Approve maupun Revoke tidak ditemukan di halaman.")
    return False, "Tombol Approve tidak ditemukan"


async def confirm_approve_modal(page: Page, dry_run: bool = False, loader: HashLoader = None) -> tuple[bool, str]:
    """
    Menunggu popup modal konfirmasi Approve (03_modalkonfirmasi.html)
    dan mengklik tombol 'Konfirmasi'.
    """
    if loader:
        loader.update("Menunggu modal konfirmasi Approve...")
    log.info("Menunggu modal konfirmasi 'Approve Assignment?' muncul...")

    # Tunggu modal muncul hingga 12 detik
    modal_dialog = page.locator("div[role='alertdialog']").first
    try:
        await modal_dialog.wait_for(state="visible", timeout=12_000)
    except Exception:
        log.error("Modal konfirmasi 'Approve Assignment?' tidak muncul!")
        return False, "Modal konfirmasi tidak muncul"

    if dry_run:
        log.info("[DRY RUN] Modal Approve terdeteksi. Melewati klik Konfirmasi & membatalkan dialog.")
        try:
            batal_btn = page.locator("div[role='alertdialog'] button:has-text('Batal')").first
            if await batal_btn.is_visible():
                await batal_btn.click()
        except Exception:
            pass
        return True, "DRY_RUN"

    # Cari tombol Konfirmasi di dalam modal
    confirm_btn = None
    for c_sel in MODAL_CONFIRM_SELECTORS:
        try:
            loc = page.locator(c_sel).first
            if await loc.count() > 0 and await loc.is_visible(timeout=2_000):
                confirm_btn = loc
                break
        except Exception:
            continue

    if not confirm_btn:
        log.error("Tombol 'Konfirmasi' di dalam modal tidak ditemukan!")
        return False, "Tombol Konfirmasi modal tidak ditemukan"

    if loader:
        loader.update("Mengonfirmasi persetujuan di modal...")
    log.info("Mengklik tombol 'Konfirmasi' pada modal Approve...")
    await confirm_btn.scroll_into_view_if_needed()
    try:
        await confirm_btn.click(timeout=4_000)
    except Exception:
        await confirm_btn.click(force=True)

    log.info("Tombol Konfirmasi berhasil diklik.")

    # Tunggu modal tertutup
    try:
        await modal_dialog.wait_for(state="hidden", timeout=12_000)
    except Exception:
        pass

    # Tunggu loading selesai
    if loader:
        loader.update("Menyimpan hasil approval ke server...")
    await wait_for_loading(page, timeout_ms=25_000, initial_delay_ms=600, loader=loader)

    # Cek pesan error / galat jika ada
    try:
        error_locs = await page.locator(".f\\:text-destructive, .tw\\:text-destructive, [role='alert']").all_inner_texts()
        clean_errors = [str(t).strip() for t in error_locs if t and str(t).strip()]
        if clean_errors:
            err_msg = " | ".join(clean_errors)[:150]
            log.warning(f"Peringatan/Galat terdeteksi setelah approve: {err_msg}")
    except Exception:
        pass

    # Buffer keamanan 700ms: pastikan server selesai memproses dan response tuntas di browser
    await page.wait_for_timeout(700)

    log.info("Assignment berhasil di-approve & tersimpan di server!")
    return True, "BERHASIL"


async def process_single_assignment(page: Page, item: dict, dry_run: bool = False, loader: HashLoader = None) -> tuple[bool, str]:
    """
    Memproses satu baris assignment:
      1. Buka URL assignment
      2. Tunggu loading halaman selesai
      3. Klik tombol Approve
      4. Konfirmasi pada modal dialog
    """
    url = item["link_fasih"]
    nama = item.get("nama_usaha") or item.get("nama_assignment") or item.get("id_assignment")

    log.info(f"Row {item['row_idx']}: {nama} ({url})")

    try:
        # 1. Buka URL assignment
        if loader:
            loader.update("Membuka URL assignment...")
        await page.goto(url, wait_until="domcontentloaded", timeout=TIMEOUT)

        # 2. Tunggu loading awal selesai
        if loader:
            loader.update("Menunggu halaman kuesioner siap...")
        await wait_for_loading(page, timeout_ms=40_000, initial_delay_ms=800, loader=loader)

        # 3. Klik tombol Approve
        ok_btn, btn_msg = await click_approve_button(page, loader=loader)
        if not ok_btn:
            return False, btn_msg

        if btn_msg == "SUDAH_APPROVED":
            return True, "SUDAH_APPROVED"

        # 4. Konfirmasi pada modal popup
        ok_confirm, confirm_msg = await confirm_approve_modal(page, dry_run=dry_run, loader=loader)
        if not ok_confirm:
            return False, confirm_msg

        return True, confirm_msg

    except Exception as e:
        log.error(f"Error saat memproses assignment: {e}")
        return False, f"Exception: {str(e)[:100]}"


# ─────────────────────────── MAIN FUNCTION ───────────────────────────

async def main():
    parser = argparse.ArgumentParser(description="FASIH Survey Automation - Approve Admin")
    parser.add_argument("--excel", type=str, default=None, help="Path ke file Excel data assignment")
    parser.add_argument("--dry-run", action="store_true", help="Simulasi tanpa klik konfirmasi akhir")
    parser.add_argument("--limit", type=int, default=None, help="Batasi jumlah data yang diproses")
    parser.add_argument("--start-row", type=int, default=None, help="Mulai dari baris Excel tertentu")
    parser.add_argument("--delay", type=float, default=0.7, help="Jeda antar assignment dalam detik (default: 0.7)")
    parser.add_argument("--all-rows", action="store_true", help="Proses semua baris terlepas dari Status Otomasi")
    args = parser.parse_args()

    print(r"""
╔══════════════════════════════════════════════════════════════════════╗
║               FASIH AUTOMATION - APPROVE ADMIN                       ║
║                                                                      ║
║  1. Pilih file Excel data assignment yang akan di-running            ║
║  2. Browser Chrome akan terbuka otomatis                             ║
║  3. Silakan LOGIN MANUAL ke akun FASIH Anda di browser (jika belum)  ║
║  4. Bot akan otomatis approve assignment satu per satu               ║
╚══════════════════════════════════════════════════════════════════════╝
    """)

    # Tentukan path file Excel (menu interaktif di terminal)
    excel_path = select_excel_interactive(cli_excel=args.excel)
    if not excel_path or not excel_path.exists():
        log.error(f"File Excel '{excel_path}' tidak valid atau tidak ditemukan!")
        return

    dry_run = args.dry_run

    wb, items = load_assignments(excel_path)

    # Muat status dari log_simple_approve.txt untuk resume otomatis
    status_map = load_log_simple_status(LOG_SIMPLE_TXT)
    log.info(f"Memuat status riwayat dari '{LOG_SIMPLE_TXT}' ({len(status_map)} entri terdaftar)...")

    pending_items = []
    sukses_logged = 0
    gagal_logged = 0
    excel_needs_save = False

    for it in items:
        # Filter berdasarkan start-row jika ada
        if args.start_row and it["row_idx"] < args.start_row:
            continue

        # Jika Excel memiliki kolom 'Status Otomasi' (seperti data_edited.xlsx hasil input),
        # dan flag --all-rows tidak aktif: hanya approve data yang Status Otomasi == 'BERHASIL'
        if not args.all_rows and it.get("status_otomasi"):
            if "BERHASIL" not in it["status_otomasi"].upper():
                continue

        id_key = _normalize_key(it.get("id_assignment") or it.get("row_idx"))
        st_log = status_map.get(id_key)

        # Cek apakah sudah pernah sukses
        is_already_success = False
        if st_log == "sukses":
            is_already_success = True
        elif any(k in str(it.get("status", "")).upper() for k in ("BERHASIL", "OK", "SUKSES", "APPROVED", "SUDAH_APPROVED")):
            is_already_success = True

        if is_already_success:
            sukses_logged += 1
            if not any(k in str(it.get("status", "")).upper() for k in ("BERHASIL", "OK", "APPROVED", "SUDAH_APPROVED")):
                it["status"] = "BERHASIL"
                ws = wb.active
                ws.cell(row=it["row_idx"], column=it["status_col"], value="BERHASIL")
                ws.cell(row=it["row_idx"], column=it["time_col"], value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                excel_needs_save = True
        elif st_log in ("tidak", "gagal"):
            gagal_logged += 1
            pending_items.append(it)
        else:
            pending_items.append(it)

    if excel_needs_save:
        try:
            wb.save(excel_path)
            log.info(f"💾 File Excel '{excel_path}' disinkronkan otomatis dengan log riwayat.")
        except Exception as e:
            log.warning(f"Tidak dapat menyimpan sinkronisasi ke Excel: {e}")

    # Batasi data jika parameter --limit diisi
    if args.limit and args.limit > 0:
        pending_items = pending_items[:args.limit]
        log.info(f"Batasan --limit aktif: hanya memproses {len(pending_items)} data pertama.")

    log.info(f"Total data di Excel       : {len(items)}")
    log.info(f"Data SUKSES di riwayat    : {sukses_logged} (di-skip)")
    if gagal_logged > 0:
        log.info(f"Data GAGAL di riwayat     : {gagal_logged} (akan dieksekusi ulang)")
    log.info(f"Total data pending/run    : {len(pending_items)}")
    if pending_items:
        nama_pertama = pending_items[0].get('nama_usaha') or pending_items[0].get('nama_assignment') or pending_items[0].get('id_assignment')
        log.info(f"Mulai eksekusi dari Baris {pending_items[0]['row_idx']}: {nama_pertama}")
    log.info(f"Mode DRY_RUN              : {dry_run}")

    if not pending_items:
        log.info("🎉 Semua assignment sudah berstatus SUKSES / APPROVED!")
        return

    # Inisialisasi storage_state jika ada
    storage_state = None
    if Path(AUTH_FILE).exists():
        storage_state = AUTH_FILE
    elif Path("input_fasih_ntb/auth.json").exists():
        storage_state = "input_fasih_ntb/auth.json"

    chrome_proc = None
    anti_sleep_handle = start_anti_sleep()
    profile_dir = Path.home() / "chrome_fasih_profile"
    profile_dir.mkdir(parents=True, exist_ok=True)

    try:
        # 1. Deteksi resolusi layar dan siapkan Chrome
        screen_w, screen_h = get_screen_bounds()
        log.info(f"🖥️ Resolusi layar terdeteksi: {screen_w} x {screen_h}. Menyesuaikan browser agar maximize...")

        chrome_exe = get_chrome_executable()
        log.info(f"Membuka Google Chrome di {BASE_URL}/ ... ({chrome_exe})")
        chrome_cmd = [
            chrome_exe,
            "--remote-debugging-port=9222",
            f"--user-data-dir={str(profile_dir)}",
            "--no-first-run",
            "--no-default-browser-check",
            "--start-maximized",
            f"--window-size={screen_w},{screen_h}",
            "--window-position=0,0",
            f"{BASE_URL}/"
        ]
        chrome_proc = subprocess.Popen(chrome_cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        await asyncio.sleep(2.5)

        async with async_playwright() as pw:
            log.info("Menghubungkan Playwright ke Google Chrome (CDP port 9222)...")
            browser = None
            for cdp_attempt in range(7):
                try:
                    browser = await pw.chromium.connect_over_cdp("http://127.0.0.1:9222")
                    break
                except Exception as ce:
                    if cdp_attempt < 6:
                        await asyncio.sleep(1.5)
                    else:
                        raise ce

            context = browser.contexts[0]
            page = context.pages[0] if context.pages else await context.new_page()
            page.set_default_timeout(TIMEOUT)

            # Maksimalkan ukuran jendela Chrome dan viewport
            await maximize_chrome_window(page, screen_w, screen_h)
            try:
                dim = await page.evaluate("() => ({ w: window.screen.availWidth, h: window.screen.availHeight })")
                vw = dim.get("w") or screen_w
                vh = dim.get("h") or screen_h
                await page.set_viewport_size({"width": vw, "height": vh})
                await maximize_chrome_window(page, vw, vh)
            except Exception:
                await page.set_viewport_size({"width": screen_w, "height": screen_h})

            # Cek apakah sudah login atau perlu login manual
            is_logged_in = False
            try:
                user_el = page.locator("span.user-name, .avatar-status-online, button[aria-label='sidebar-toggle'], div.fasih-form-sidebar").first
                if await user_el.is_visible(timeout=5_000):
                    is_logged_in = True
                    log.info("✅ Sudah dalam keadaan login!")
            except Exception:
                pass

            if not is_logged_in:
                print("\n" + "=" * 60)
                print("⏳ GOOGLE CHROME TERBUKA DI https://fasih-sm.bps.go.id/")
                print("Silakan LOGIN MANUAL seperti biasa.")
                print("Setelah login berhasil dan masuk ke FASIH,")
                print("TEKAN ENTER DI TERMINAL INI UNTUK MEMULAI OTOMASI APPROVE ADMIN...")
                print("=" * 60)
                input()

            # Simpan state login untuk penggunaan berikutnya
            try:
                await context.storage_state(path=AUTH_FILE)
                log.info(f"Session disimpan ke {AUTH_FILE}")
            except Exception:
                pass

            # Mulai perulangan setiap assignment
            sukses = 0
            gagal = 0
            approve_baru_count = 0
            sudah_approved_count = 0
            total_pending = len(pending_items)

            print(f"\n🚀 Memulai pemrosesan {total_pending} assignment...")
            print("─" * 70)

            for i, item in enumerate(pending_items, 1):
                nama_item = item.get("nama_usaha") or item.get("nama_assignment") or item.get("id_assignment")
                p_str = format_progress(i, total_pending)
                log.info(f"\n[{p_str}] Memproses Baris {item['row_idx']} di Excel: {nama_item}...")

                update_last_row_txt(
                    item,
                    "SEDANG DIPROSES...",
                    current_idx=i,
                    total_items=total_pending,
                    sukses=sukses,
                    gagal=gagal,
                )

                # Animasi loading terminal '####' dengan format persentase dan baris: [#######---] 70% (7/10 baris)
                loader = HashLoader(
                    message="Membuka assignment...",
                    prefix=f"\033[94m{p_str}\033[0m |"
                )
                loader.start()

                # Watchdog timer per baris
                try:
                    ok, status_msg = await asyncio.wait_for(
                        process_single_assignment(page, item, dry_run=dry_run, loader=loader),
                        timeout=ASSIGNMENT_ROW_TIMEOUT
                    )
                except asyncio.TimeoutError:
                    ok = False
                    status_msg = f"Stuck/Timeout: Melebihi batas waktu maksimal ({ASSIGNMENT_ROW_TIMEOUT}s)"
                    log.error(
                        f"❌ [TIMEOUT / STUCK] Baris {item['row_idx']} ({nama_item}) stuck > {ASSIGNMENT_ROW_TIMEOUT} detik! "
                        f"Lanjut ke baris berikutnya..."
                    )
                finally:
                    await loader.stop(clear=True)

                if ok:
                    sukses += 1
                    if status_msg == "SUDAH_APPROVED":
                        sudah_approved_count += 1
                        save_val = "BERHASIL (SUDAH_APPROVED)"
                        log.info(f"⏭️ Baris {item['row_idx']} DI-SKIP: Sudah approved admin sebelumnya (Tombol Revoke terdeteksi, 04_revoke.html). Status Excel: {save_val}.")
                        note_text = "Sudah approved admin sebelumnya (Tombol Revoke terdeteksi) - di-skip"
                        print(f"\033[94m{p_str}\033[0m | Baris {item['row_idx']} ({nama_item}): \033[93m⏭️  SUDAH APPROVED\033[0m (di-skip)")
                    else:
                        approve_baru_count += 1
                        save_val = "BERHASIL"
                        log.info(f"✅ Baris {item['row_idx']} SELESAI ({status_msg}). Status Excel di-update: {save_val}.")
                        note_text = status_msg
                        print(f"\033[94m{p_str}\033[0m | Baris {item['row_idx']} ({nama_item}): \033[92m✅ BERHASIL DI-APPROVE\033[0m")

                    save_assignment_status(wb, excel_path, item, save_val)
                    update_last_row_txt(
                        item,
                        save_val,
                        note=note_text,
                        current_idx=i,
                        total_items=total_pending,
                        sukses=sukses,
                        gagal=gagal,
                    )
                    append_log_simple(item, "sukses")
                else:
                    gagal += 1
                    save_assignment_status(wb, excel_path, item, "GAGAL", note=status_msg)
                    log.warning(
                        f"⚠️ Baris {item['row_idx']} di Excel ({nama_item}) GAGAL / DI-SKIP!\n"
                        f"   Alasan: {status_msg}\n"
                        f"   Status disimpan sebagai GAGAL agar otomatis dieksekusi ulang saat re-run."
                    )
                    print(f"\033[94m{p_str}\033[0m | Baris {item['row_idx']} ({nama_item}): \033[91m❌ GAGAL\033[0m ({status_msg})")
                    update_last_row_txt(
                        item,
                        "GAGAL",
                        note=f"{status_msg} (Pending di Excel untuk re-run)",
                        current_idx=i,
                        total_items=total_pending,
                        sukses=sukses,
                        gagal=gagal,
                    )
                    append_log_simple(item, "tidak")

                if args.delay > 0:
                    await asyncio.sleep(args.delay)

            print("\n" + "=" * 60)
            print("RINGKASAN EKSEKUSI APPROVE ADMIN:")
            print(f"Total diproses           : {len(pending_items)}")
            print(f"✓ Berhasil Approve baru  : {approve_baru_count}")
            print(f"⏭ Sudah Approved (Skip)  : {sudah_approved_count}")
            print(f"✗ Gagal                  : {gagal}")
            print(f"File Excel               : {excel_path} (status terupdate)")
            print(f"File Log                 : {LOG_FILE}")
            print(f"Log Ringkas              : {LOG_SIMPLE_TXT}")
            print(f"Log Terakhir             : {LAST_ROW_TXT}")
            print("=" * 60)

            input("\nTekan ENTER untuk selesai...")
            await browser.close()

    finally:
        if chrome_proc:
            try:
                if sys.platform == "win32":
                    subprocess.run(["taskkill", "/F", "/T", "/PID", str(chrome_proc.pid)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                else:
                    chrome_proc.terminate()
            except Exception:
                pass
        stop_anti_sleep(anti_sleep_handle)


if __name__ == "__main__":
    asyncio.run(main())
